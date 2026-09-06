"""Export the Markdown research report to DOCX."""

import re

from docx import Document

from analysis_support import REPORTS_DIR


def main() -> None:
    md_path = REPORTS_DIR / "REPORT.md"
    docx_path = REPORTS_DIR / "REPORT.docx"

    text = md_path.read_text(encoding="utf-8")
    doc = Document()

    for raw in text.splitlines():
        line = raw.rstrip()
        if not line:
            doc.add_paragraph("")
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
            continue

        if line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
            continue

        numbered_item = re.match(r"\d+\.\s+(.*)", line)
        if numbered_item:
            doc.add_paragraph(numbered_item.group(1), style="List Number")
            continue

        # Markdown tables remain plain text; richer conversion is outside this exporter.
        doc.add_paragraph(line)

    doc.save(docx_path)
    print(f"DOCX generated: {docx_path}")


if __name__ == "__main__":
    main()

